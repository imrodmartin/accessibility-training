from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import copy

NAVY = RGBColor(0x15,0x0B,0x4F)
ORANGE = RGBColor(0xFF,0x85,0x00)
DARK = RGBColor(0x1A,0x1A,0x1A)

SLIDES = [
    ("Mill Street water main replacement",
     ["Council briefing", "Department of Public Works", "March 20, 2026"]),
    ("Why the project is needed",
     ["The main under Mill Street dates to 1954.",
      "Three breaks since 2023, two of them in winter.",
      "Repair cost has passed 40 percent of replacement cost."]),
    ("Scope of work",
     ["Roughly 1,800 feet, Water Street to Rail Avenue.",
      "New 12-inch ductile iron main and 41 service connections.",
      "Full-depth pavement restoration curb to curb."]),
    ("Schedule and impact",
     ["Construction runs April 6 to May 15.",
      "One lane, alternating, 7:00 a.m. to 5:00 p.m. weekdays.",
      "Service interruptions up to six hours, 48 hours notice."]),
    ("What we need from Council",
     ["Approve the construction contract at $1.84 million.",
      "Authorize the detour on Water Street through May 15.",
      "Questions to Public Works at (555) 555-0142."]),
]

def add_box(slide, l, t, w, h, text, size, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(text if isinstance(text, list) else [text]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
        r.font.name = "Calibri"
        p.alignment = align
        p.space_after = Pt(10)
    return tb

def set_alt(shape, title, desc):
    el = shape._element._nvXxPr.cNvPr
    el.set("title", title)
    el.set("descr", desc)

# ================= BEFORE =================
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]          # Defect: blank layout, no title placeholder anywhere

for i, (title, bullets) in enumerate(SLIDES):
    s = prs.slides.add_slide(blank)
    # Defect: bullets placed FIRST, so the reading order hits the body before the title
    add_box(s, 1.0, 2.2, 11.0, 4.0, ["- " + b for b in bullets], 20)
    # Defect: the "title" is a floating text box that only looks like a title
    add_box(s, 1.0, 0.9, 11.0, 1.0, title, 34, bold=True, color=NAVY)
    # Defect: decorative rule as a shape with no alt text handling
    line = s.shapes.add_textbox(Inches(1.0), Inches(1.75), Inches(3.0), Inches(0.1))
    line.fill.solid(); line.fill.fore_color.rgb = ORANGE
    # Defect: page number as an unlabelled floating box
    add_box(s, 12.4, 6.9, 0.6, 0.4, str(i+1), 12)

# Defect: an image with no alternative text
s2 = prs.slides[2]
pic = s2.shapes.add_picture("city-seal.png", Inches(11.0), Inches(0.5), height=Inches(1.1))
for a in ("descr","title"):
    if pic._element._nvXxPr.cNvPr.get(a) is not None:
        del pic._element._nvXxPr.cNvPr.attrib[a]

prs.core_properties.title = ""        # Defect: no deck title
prs.core_properties.author = "jgreene"
prs.save("EX3-deck-BEFORE.pptx")

# ================= AFTER =================
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
title_layout   = prs.slide_layouts[0]   # Title Slide
content_layout = prs.slide_layouts[1]   # Title and Content

for i, (title, bullets) in enumerate(SLIDES):
    if i == 0:
        s = prs.slides.add_slide(title_layout)
        s.shapes.title.text = title                 # Fix: real title placeholder
        sub = s.placeholders[1].text_frame
        sub.text = bullets[0]
        for b in bullets[1:]:
            sub.add_paragraph().text = b
    else:
        s = prs.slides.add_slide(content_layout)
        s.shapes.title.text = title                 # Fix: unique title on every slide
        body = s.placeholders[1].text_frame         # Fix: content placeholder, real list
        body.text = bullets[0]
        for b in bullets[1:]:
            body.add_paragraph().text = b
        for p in body.paragraphs:
            for r in p.runs:
                r.font.size = Pt(20)
    for p in s.shapes.title.text_frame.paragraphs:
        for r in p.runs:
            r.font.color.rgb = NAVY

# Fix: the image carries alternative text, and is placed after the content
s2 = prs.slides[2]
pic = s2.shapes.add_picture("city-seal.png", Inches(11.4), Inches(0.4), height=Inches(1.1))
set_alt(pic, "City of Riverton seal",
        "Seal of the City of Riverton, incorporated 1887, Office of the City Clerk.")

prs.core_properties.title = "Mill Street water main replacement, Council briefing"
prs.core_properties.author = "City of Riverton, Department of Public Works"
prs.core_properties.language = "en-US"
prs.core_properties.comments = "Five slides. Every slide uses the title placeholder; content sits in placeholders, not floating text boxes."
prs.save("EX3-deck-AFTER.pptx")
print("decks written")
