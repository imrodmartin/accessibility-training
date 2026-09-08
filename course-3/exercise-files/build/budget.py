from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROWS = [
 ("Public Works","Mill Street water main", 1_840_000, 1_612_400, "on track"),
 ("Public Works","Fleet replacement",        420_000,   455_900, "over"),
 ("Public Works","Snow and ice control",     310_000,   298_100, "on track"),
 ("Parks","Riverfront trail resurfacing",    265_000,   119_750, "delayed"),
 ("Parks","Playground safety surfacing",      88_000,    86_400, "on track"),
 ("Library","Collection and databases",      196_000,   190_200, "on track"),
 ("Library","HVAC replacement",              240_000,   271_300, "over"),
 ("City Clerk","Records digitization",       132_000,    54_600, "delayed"),
 ("City Clerk","Election administration",     97_500,    95_800, "on track"),
 ("Facilities","Badge system upgrade",       145_000,   141_200, "on track"),
]
RED    = "FFC7CE"   # over
YELLOW = "FFEB9C"   # delayed
GREEN  = "C6EFCE"   # on track
FILLS = {"over":RED, "delayed":YELLOW, "on track":GREEN}

def money(ws, col, first, last):
    for r in range(first, last+1):
        ws.cell(row=r, column=col).number_format = '"$"#,##0'

# ---------------- BEFORE ----------------
wb = Workbook()
ws = wb.active                      # Defect: sheet left as "Sheet1"
ws.title = "Sheet1"
wb.create_sheet("Sheet2"); wb.create_sheet("Sheet3")   # Defect: empty generic sheets

# Defect: title floating in a merged banner, not a real header structure
ws.merge_cells("A1:E1")
ws["A1"] = "FY2026 Q2 BUDGET STATUS"
ws["A1"].font = Font(bold=True, size=16)
ws["A1"].alignment = Alignment(horizontal="center")
# Defect: blank spacer row
r = 3
# Defect: NO header row at all - data starts immediately
for dept, item, budget, spent, status in ROWS:
    ws.cell(row=r, column=1, value=dept)
    ws.cell(row=r, column=2, value=item)
    ws.cell(row=r, column=3, value=budget)
    ws.cell(row=r, column=4, value=spent)
    # Defect: status is colour only, no text
    ws.cell(row=r, column=5).fill = PatternFill("solid", fgColor=FILLS[status])
    r += 1
money(ws, 3, 3, r-1); money(ws, 4, 3, r-1)
# Defect: legend as a floating note far from the data
ws["G3"] = "red = over, yellow = delayed, green = on track"
ws["G3"].font = Font(italic=True)
for c,w in zip("ABCDEFG",[16,34,14,14,6,3,42]): ws.column_dimensions[c].width = w
wb.properties.title = None          # Defect: no document title
wb.properties.creator = "jgreene"
wb.save("EX3-budget-BEFORE.xlsx")

# ---------------- AFTER ----------------
wb = Workbook()
ws = wb.active
ws.title = "Q2 budget status"       # Fix: descriptive sheet name, no empty sheets

hdr = ["Department","Line item","Budget","Spent to date","Status"]
ws.append(hdr)                      # Fix: real header row, row 1, no spacer
head_fill = PatternFill("solid", fgColor="150B4F")
thin = Side(style="thin", color="9A98A8")
for i in range(1, len(hdr)+1):
    c = ws.cell(row=1, column=i)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = head_fill
    c.alignment = Alignment(vertical="center")
    c.border = Border(bottom=thin)

for dept, item, budget, spent, status in ROWS:
    # Fix: status is a text value; colour reinforces, never carries, the meaning
    ws.append([dept, item, budget, spent, status.title()])
    row = ws.max_row
    ws.cell(row=row, column=5).fill = PatternFill("solid", fgColor=FILLS[status])

last = ws.max_row
money(ws, 3, 2, last); money(ws, 4, 2, last)

# Fix: a defined table gives the header row a programmatic relationship to the data
t = Table(displayName="Q2BudgetStatus", ref=f"A1:E{last}")
t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
ws.add_table(t)
ws.freeze_panes = "A2"              # Fix: header stays visible
ws.auto_filter.ref = f"A1:E{last}"
for c,w in zip("ABCDE",[16,34,14,16,12]): ws.column_dimensions[c].width = w

ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = "1:1"         # Fix: header repeats when printed
wb.properties.title = "FY2026 Q2 budget status, City of Riverton"
wb.properties.creator = "City of Riverton, Finance Department"
wb.properties.language = "en-US"
wb.properties.description = "Quarterly budget status by department. Status is stated as text and reinforced with color."
wb.save("EX3-budget-AFTER.xlsx")
print("budgets written")
